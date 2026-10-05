"""Assemble a GOST-profile DOCX from a generated semantic body and a cover template.

The body is an editable DOCX (for example, generated with python-docx or Pandoc),
not a PDF conversion. Render the assembled document with finalize_docx.py.
"""
from __future__ import annotations
import argparse
import hashlib
import json
from pathlib import Path
import re
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml.ns import qn
from docxcompose.composer import Composer
from lxml import etree

from gost_docx import add_toc, format_body, number_equations, page_footer


def fill_cover(template, values):
    """Replace {{field}} across split runs without rewriting paragraph formatting."""
    document = Document(template)
    for paragraph in document.element.findall('.//' + qn('w:p')):
        nodes = paragraph.findall('.//' + qn('w:t'))
        text = ''.join(node.text or '' for node in nodes)
        for match in reversed(list(re.finditer(r'\{\{([a-z_]+)\}\}', text))):
            key = match.group(1)
            if key not in values:
                raise ValueError('Missing cover field: ' + key)
            start, end = match.span()
            offset = 0
            inserted = False
            for node in nodes:
                old = node.text or ''
                a, b = offset, offset + len(old)
                offset = b
                if b <= start or a >= end:
                    continue
                left, right = max(0, start - a), min(len(old), end - a)
                node.text = old[:left] + (values[key] if not inserted else '') + old[right:]
                inserted = True
                node.set(qn('xml:space'), 'preserve')
    # Markers in unsupported parts (including headers) must not survive silently.
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    for part in document.part.package.parts:
        if not part.content_type.endswith('xml'):
            continue
        tree = etree.fromstring(part.blob, parser)
        for paragraph in tree.findall('.//' + qn('w:p')):
            text = ''.join(node.text or '' for node in paragraph.findall('.//' + qn('w:t')))
            if '{{' in text:
                raise ValueError('Unresolved cover marker; use body paragraphs or layout-table cells')
    return document


def assemble(body, cover, values_path, output):
    if output.exists():
        raise ValueError('Use a new output path; source files must not be overwritten')
    values = json.loads(values_path.read_text())
    if not isinstance(values, dict) or not all(isinstance(k, str) and isinstance(v, str) for k, v in values.items()):
        raise ValueError('Cover fields must be a JSON object of string values')
    family = subprocess.run(['fc-match', '--format=%{family}', 'Times New Roman'],
                            capture_output=True, text=True, check=True).stdout
    if 'Times New Roman' not in family.split(','):
        raise ValueError('Install Times New Roman legally; no silent font substitution')
    document = Document(body)
    if any('TOC ' in (node.text or '') for node in document.element.findall('.//' + qn('w:instrText'))):
        raise ValueError('Supply a body without an existing TOC; the assembly creates its own')
    format_body(document)
    formulas = [p for p in document.paragraphs if p._p.find(qn('m:oMathPara')) is not None]
    if formulas and any(table._tbl.find('.//' + qn('m:oMath')) is not None for table in document.tables):
        raise ValueError('Mixed pre-numbered and unnumbered equation layouts require explicit source preparation')
    number_equations(document, list(range(1, len(formulas) + 1)))
    add_toc(document)
    filled = fill_cover(cover, values)
    if len(filled.sections) != 1:
        raise ValueError('The cover template must contain one section; inspect complex covers separately')
    first = filled.sections[0]
    first.different_first_page_header_footer = True
    if any('PAGE' in (node.text or '') for node in first.first_page_footer._element.findall('.//' + qn('w:instrText'))):
        raise ValueError('The title-page footer must not contain a visible PAGE field')
    section = filled.add_section(WD_SECTION_START.NEW_PAGE)
    source_section = document.sections[-1]
    for name in ('page_width', 'page_height', 'left_margin', 'right_margin', 'top_margin', 'bottom_margin', 'footer_distance'):
        setattr(section, name, getattr(source_section, name))
    section.different_first_page_header_footer = False
    section.footer.is_linked_to_previous = False
    page_footer(section, '2')
    composer = Composer(filled)
    composer.append(document)
    output.parent.mkdir(parents=True, exist_ok=True)
    composer.save(output)
    manifest = {
        'route': 'editable DOCX master; PDF exported from the finalized DOCX',
        'inputs': {str(path): hashlib.sha256(path.read_bytes()).hexdigest() for path in (body, cover, values_path)},
        'docx_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
        'cover': 'field replacements preserve runs/layout; original template retained',
        'review_required': ['Finalize TOC/PAGE fields', 'Render and inspect every page',
                            'Check cover occupies one page and has no visible number',
                            'Verify continuation captions on every continued table part',
                            'Check all content, equation numbers and bibliography against sources'],
    }
    (output.parent / 'docx-manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
    print(f'Assembled {output}; finalize fields and inspect its PDF before delivery.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--body', required=True, type=Path)
    parser.add_argument('--cover', required=True, type=Path)
    parser.add_argument('--fields', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    try:
        assemble(args.body.resolve(), args.cover.resolve(), args.fields.resolve(), args.output.resolve())
    except (OSError, ValueError, subprocess.CalledProcessError) as error:
        parser.exit(1, str(error) + '\n')


if __name__ == '__main__':
    main()
