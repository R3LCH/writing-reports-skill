"""Shared student-report Word profile; content and institutional covers are separate."""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import re
from copy import deepcopy

STRUCTURAL = {
    'СПИСОК ИСПОЛНИТЕЛЕЙ', 'РЕФЕРАТ', 'СОДЕРЖАНИЕ', 'ТЕРМИНЫ И ОПРЕДЕЛЕНИЯ',
    'ПЕРЕЧЕНЬ СОКРАЩЕНИЙ И ОБОЗНАЧЕНИЙ', 'ВВЕДЕНИЕ', 'ЗАКЛЮЧЕНИЕ',
    'СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ',
}


def field(paragraph, instruction, result=''):
    for kind, text in [('begin', None), (None, instruction), ('separate', None)]:
        node = OxmlElement('w:instrText' if kind is None else 'w:fldChar')
        if kind is None:
            node.text = text
            node.set(qn('xml:space'), 'preserve')
        else:
            node.set(qn('w:fldCharType'), kind)
        paragraph.add_run()._r.append(node)
    paragraph.add_run(result)
    node = OxmlElement('w:fldChar')
    node.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.append(node)


def fonts(properties, name):
    node = properties.find(qn('w:rFonts'))
    if node is None:
        node = OxmlElement('w:rFonts')
        properties.insert(0, node)
    for attribute in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        node.set(qn('w:' + attribute), name)
    for attribute in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        node.attrib.pop(qn('w:' + attribute), None)


def run_font(run, *, size=14, mono=False):
    name = 'DejaVu Sans Mono' if mono else 'Times New Roman'
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts(run._r.get_or_add_rPr(), name)
    if mono:
        run.bold = False
        run.italic = False


def page_footer(section, result):
    footer = section.footer.paragraphs[0]
    footer.clear()
    fmt = footer.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt.first_line_indent = Cm(0)
    fmt.left_indent = Cm(0)
    # Center on the physical page, not the asymmetric 30/15 mm text block.
    fmt.right_indent = section.left_margin - section.right_margin
    field(footer, ' PAGE ', result)
    for run in footer.runs:
        run_font(run)
        run.bold = False
        run.italic = False


def table_borders(table, value):
    properties = table._tbl.tblPr
    old = properties.find(qn('w:tblBorders'))
    if old is not None:
        properties.remove(old)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + side)
        border.set(qn('w:val'), value)
        if value != 'nil':
            border.set(qn('w:sz'), '4')
            border.set(qn('w:color'), '000000')
        borders.append(border)
    properties.append(borders)


def appendix_titles(document):
    """Keep label/title on separate lines and include the full name in the TOC."""
    paragraphs = list(document.paragraphs)
    for index, heading in enumerate(paragraphs):
        if not heading.style.name.startswith('Heading'):
            continue
        if not re.fullmatch(r'ПРИЛОЖЕНИЕ [А-ЯA-Z0-9]+(?: \([^)]+\))?', heading.text.upper()):
            continue
        if index + 1 == len(paragraphs):
            raise ValueError('An appendix label needs its title on the next paragraph')
        title = paragraphs[index + 1]
        if not title.text.strip() or heading._p.getnext() is not title._p:
            raise ValueError('Place the appendix title immediately after its label')
        for run in heading.runs:
            run.text = run.text.upper()
        heading.add_run().add_break()
        for child in title._p:
            if child.tag != qn('w:pPr'):
                heading._p.append(deepcopy(child))
        title._p.getparent().remove(title._p)


def format_body(document, *, structural=(), includes_cover=False):
    """Apply to generated body only; never to an institutional cover template."""
    structural_names = STRUCTURAL | {name.upper() for name in structural}
    appendix_titles(document)
    for section in document.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.left_margin, section.right_margin = Cm(3), Cm(1.5)
        section.top_margin, section.bottom_margin = Cm(2), Cm(2)
        section.footer_distance = Cm(1)
        section.different_first_page_header_footer = includes_cover
        page_footer(section, '2' if includes_cover else '1')
    for style in document.styles:
        if style.type not in (1, 2):
            continue
        mono = style.name in ('Source Code', 'Verbatim Char')
        name = 'DejaVu Sans Mono' if mono else 'Times New Roman'
        style.font.name = name
        style.font.size = Pt(14)
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts(style.element.get_or_add_rPr(), name)
        if style.type == 1:
            style.font.bold = False
            style.font.italic = False
            paragraph = style.paragraph_format
            paragraph.line_spacing = 1.5
            paragraph.space_before = Pt(0)
            paragraph.space_after = Pt(0)
            paragraph.first_line_indent = Cm(1.25)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    cover = includes_cover
    cover_paragraphs = []
    for paragraph in document.paragraphs:
        fmt = paragraph.paragraph_format
        fmt.line_spacing = 1.5
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.first_line_indent = Cm(1.25)
        fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if paragraph.style.name.startswith('Heading'):
            cover = False
            level = int(paragraph.style.name.split()[-1])
            text = paragraph.text.upper()
            is_appendix = re.match(r'ПРИЛОЖЕНИЕ [А-ЯA-Z0-9]+(?:\s|$)', text)
            is_structural = text in structural_names or is_appendix
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_structural else WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Cm(0 if is_structural else 1.25)
            fmt.keep_with_next = True
            fmt.page_break_before = level == 1
            properties = paragraph._p.get_or_add_pPr()
            if properties.find(qn('w:suppressAutoHyphens')) is None:
                properties.append(OxmlElement('w:suppressAutoHyphens'))
            for run in paragraph.runs:
                if is_structural and not is_appendix:
                    run.text = run.text.upper()
                run.bold = True
        elif cover:
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent = Cm(0)
            cover_paragraphs.append(paragraph)
        elif paragraph.text.startswith(('Таблица ', 'Продолжение таблицы ')):
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Cm(0)
            fmt.line_spacing = 1
            fmt.keep_with_next = True
        elif paragraph.text.startswith('Рисунок '):
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent = Cm(0)
        elif paragraph.style.name == 'Source Code':
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            fmt.first_line_indent = Cm(0)
        elif paragraph.text.startswith('где '):
            fmt.first_line_indent = Cm(0)
        if paragraph._p.find('.//' + qn('w:drawing')) is not None:
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.first_line_indent = Cm(0)
            fmt.keep_with_next = True
        for run in paragraph.runs:
            run_font(run, mono=paragraph.style.name == 'Source Code' or run.style.name == 'Verbatim Char')
    # Preserve equation layout tables; unlike data tables they have no borders.
    for table in document.tables:
        if table._tbl.find('.//' + qn('m:oMath')) is not None:
            continue
        table.autofit = False
        widths = [column.width or Cm(1) for column in table.columns]
        total = sum(widths)
        for column, width in zip(table.columns, widths):
            column.width = int(Cm(16.5) * width / total)
        table_borders(table, 'single')
        for index, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn('w:cantSplit')) is None:
                properties.append(OxmlElement('w:cantSplit'))
            for col, cell in enumerate(row.cells):
                cell.width = table.columns[col].width
                for paragraph in cell.paragraphs:
                    fmt = paragraph.paragraph_format
                    fmt.first_line_indent = Cm(0)
                    fmt.space_before = Pt(0)
                    fmt.space_after = Pt(0)
                    fmt.line_spacing = 1
                    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or col > 0 else WD_ALIGN_PARAGRAPH.LEFT
                    for run in paragraph.runs:
                        run_font(run, size=12)
                        run.bold = False
    return cover_paragraphs


def number_equations(document, numbers):
    formulas = [p for p in document.paragraphs if p._p.find(qn('m:oMathPara')) is not None]
    if len(formulas) != len(numbers):
        raise ValueError('Display equations changed; inspect the semantic source')
    for paragraph, number in zip(formulas, numbers):
        if number is None:
            continue
        table = document.add_table(rows=1, cols=3)
        table.autofit = False
        for column, width in zip(table.columns, (2, 12.5, 2)):
            column.width = Cm(width)
        for cell, width in zip(table.rows[0].cells, (2, 12.5, 2)):
            cell.width = Cm(width)
            fmt = cell.paragraphs[0].paragraph_format
            fmt.first_line_indent = Cm(0)
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(0)
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        table.cell(0, 1).paragraphs[0]._p.append(paragraph._p.find(qn('m:oMathPara')))
        label = table.cell(0, 2).paragraphs[0]
        label.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run_font(label.add_run(f'({number})'))
        table_borders(table, 'nil')
        before = paragraph.insert_paragraph_before()
        before.paragraph_format.keep_with_next = True
        before.paragraph_format.first_line_indent = Cm(0)
        paragraph._p.addprevious(table._tbl)
        paragraph.clear()
        paragraph.paragraph_format.first_line_indent = Cm(0)


def add_toc(document):
    first = next((p for p in document.paragraphs if p.style.name.startswith('Heading')), None)
    if first is None:
        raise ValueError('No report headings: cannot establish the TOC boundary')
    title = first.insert_paragraph_before('СОДЕРЖАНИЕ')
    title.style = document.styles['TOC Heading']
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Cm(0)
    run_font(title.runs[0])
    title.runs[0].bold = True
    toc = first.insert_paragraph_before()
    toc.paragraph_format.first_line_indent = Cm(0)
    field(toc, ' TOC \\o "1-3" \\h \\z \\u ', 'Обновите оглавление перед сдачей.')
    if document.settings.element.find(qn('w:updateFields')) is None:
        setting = OxmlElement('w:updateFields')
        setting.set(qn('w:val'), 'true')
        document.settings.element.append(setting)
    return title
